from __future__ import annotations

import re
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from contractlab.legal_ai.schemas import ContractDraft

PRESET_NAME = 'standard_business_brief'
HEADER_PATTERN = 'memo_masthead'
CONTENT_WIDTH_DXA = 9360


def _set_font(run, *, size: float, color: str = '000000', bold: bool = False) -> None:
    run.font.name = 'Calibri'
    run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'), 'Calibri')
    run._element.get_or_add_rPr().rFonts.set(qn('w:hAnsi'), 'Calibri')
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def _set_spacing(style, *, before: float, after: float, line_spacing: float) -> None:
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def _configure_styles(document: Document) -> None:
    normal = document.styles['Normal']
    normal.font.name = 'Calibri'
    normal._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
    normal._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
    normal.font.size = Pt(11)
    _set_spacing(normal, before=0, after=6, line_spacing=1.10)

    tokens = {
        'Heading 1': (16, '2E74B5', 16, 8),
        'Heading 2': (13, '2E74B5', 12, 6),
        'Heading 3': (12, '1F4D78', 8, 4),
    }
    for name, (size, color, before, after) in tokens.items():
        style = document.styles[name]
        style.font.name = 'Calibri'
        style._element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        style._element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        _set_spacing(style, before=before, after=after, line_spacing=1.0)
        style.paragraph_format.keep_with_next = True


def _paragraph_bottom_border(paragraph, *, color: str = '2E74B5', size: int = 12) -> None:
    properties = paragraph._p.get_or_add_pPr()
    borders = properties.find(qn('w:pBdr'))
    if borders is None:
        borders = OxmlElement('w:pBdr')
        properties.append(borders)
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), str(size))
    bottom.set(qn('w:space'), '6')
    bottom.set(qn('w:color'), color)
    borders.append(bottom)


def _add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run('Page ')
    _set_font(run, size=9, color='666666')
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instruction = OxmlElement('w:instrText')
    instruction.set(qn('xml:space'), 'preserve')
    instruction.text = ' PAGE '
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    run._r.extend([begin, instruction, end])


def _add_numbering(document: Document) -> int:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn('w:abstractNumId')))
        for node in numbering.findall(qn('w:abstractNum'))
    ]
    num_ids = [int(node.get(qn('w:numId'))) for node in numbering.findall(qn('w:num'))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    level = OxmlElement('w:lvl')
    level.set(qn('w:ilvl'), '0')
    start = OxmlElement('w:start')
    start.set(qn('w:val'), '1')
    num_format = OxmlElement('w:numFmt')
    num_format.set(qn('w:val'), 'decimal')
    level_text = OxmlElement('w:lvlText')
    level_text.set(qn('w:val'), '%1.')
    suffix = OxmlElement('w:suff')
    suffix.set(qn('w:val'), 'space')
    p_style = OxmlElement('w:pStyle')
    p_style.set(qn('w:val'), 'Heading1')
    p_properties = OxmlElement('w:pPr')
    tabs = OxmlElement('w:tabs')
    tab = OxmlElement('w:tab')
    tab.set(qn('w:val'), 'num')
    tab.set(qn('w:pos'), '720')
    tabs.append(tab)
    indent = OxmlElement('w:ind')
    indent.set(qn('w:left'), '720')
    indent.set(qn('w:hanging'), '360')
    p_properties.extend([tabs, indent])
    level.extend([start, num_format, level_text, suffix, p_style, p_properties])
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement('w:num')
    num.set(qn('w:numId'), str(num_id))
    abstract_ref = OxmlElement('w:abstractNumId')
    abstract_ref.set(qn('w:val'), str(abstract_id))
    num.append(abstract_ref)
    numbering.append(num)
    return num_id


def _apply_numbering(paragraph, num_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    num_properties = OxmlElement('w:numPr')
    level = OxmlElement('w:ilvl')
    level.set(qn('w:val'), '0')
    number = OxmlElement('w:numId')
    number.set(qn('w:val'), str(num_id))
    num_properties.extend([level, number])
    properties.append(num_properties)


def _add_labeled_paragraph(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(3)
    label_run = paragraph.add_run(label + ': ')
    _set_font(label_run, size=10, color='555555', bold=True)
    value_run = paragraph.add_run(value)
    _set_font(value_run, size=10, color='222222')


def _scrub_metadata(data: bytes) -> bytes:
    source = BytesIO(data)
    target = BytesIO()
    with ZipFile(source) as archive, ZipFile(target, 'w', ZIP_DEFLATED) as output:
        for item in archive.infolist():
            if item.filename == 'docProps/custom.xml':
                continue
            payload = archive.read(item.filename)
            if item.filename.startswith('word/') and item.filename.endswith('.xml'):
                payload = re.sub(rb'\s+w:rsid[A-Za-z]*="[^"]*"', b'', payload)
            output.writestr(item, payload)
    return target.getvalue()


def build_contract_docx(draft: ContractDraft) -> bytes:
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_styles(document)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run('AI-ENABLED CONTRACTS LAB  |  DRAFT FOR REVIEW')
    _set_font(run, size=8.5, color='666666', bold=True)
    _paragraph_bottom_border(header, color='D7DBE2', size=6)
    _add_page_number(section.footer.paragraphs[0])

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    _set_font(kicker.add_run('DRAFT CONTRACT'), size=9, color='7A5A00', bold=True)
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(4)
    _set_font(title.add_run(draft.title), size=24, color='0B2545', bold=True)
    status = document.add_paragraph()
    status.paragraph_format.space_after = Pt(10)
    _set_font(status.add_run('DRAFT FOR REVIEW - NOT EXECUTION READY'), size=10, color='9B1C1C', bold=True)
    _add_labeled_paragraph(document, 'Contract type', draft.contract_type)
    _add_labeled_paragraph(document, 'Jurisdiction assumption', draft.jurisdiction or '[JURISDICTION TO BE CONFIRMED]')
    if draft.effective_date:
        _add_labeled_paragraph(document, 'Effective date', draft.effective_date)
    if draft.parties:
        _add_labeled_paragraph(document, 'Parties', '; '.join(draft.parties))
    _paragraph_bottom_border(document.paragraphs[-1], color='2E74B5', size=12)

    purpose = document.add_paragraph()
    purpose.paragraph_format.space_before = Pt(12)
    purpose.add_run(draft.purpose)

    numbering_id = _add_numbering(document)
    for draft_section in draft.sections:
        heading = document.add_paragraph(draft_section.heading, style='Heading 1')
        _apply_numbering(heading, numbering_id)
        for body in draft_section.paragraphs:
            paragraph = document.add_paragraph(body)
            paragraph.paragraph_format.keep_together = True

    if draft.unresolved_terms:
        document.add_paragraph('Terms requiring party or counsel input', style='Heading 2')
        for index, item in enumerate(draft.unresolved_terms, start=1):
            _add_labeled_paragraph(document, f'Open term {index}', item)
    if draft.assumptions:
        document.add_paragraph('Drafting assumptions', style='Heading 2')
        for index, item in enumerate(draft.assumptions, start=1):
            _add_labeled_paragraph(document, f'Assumption {index}', item)

    document.add_paragraph('Drafting note', style='Heading 2')
    note = document.add_paragraph(draft.drafting_note)
    note.paragraph_format.keep_together = True

    properties = document.core_properties
    properties.title = draft.title
    properties.subject = 'Contract draft generated for human review'
    properties.author = ''
    properties.last_modified_by = ''
    properties.keywords = ''
    buffer = BytesIO()
    document.save(buffer)
    return _scrub_metadata(buffer.getvalue())


def contract_text_bytes(draft: ContractDraft) -> bytes:
    return draft.as_text().encode('utf-8')


def safe_contract_filename(title: str) -> str:
    stem = re.sub(r'[^A-Za-z0-9._-]+', '-', title).strip('-._').lower()
    return stem[:80] or 'contract-draft'
