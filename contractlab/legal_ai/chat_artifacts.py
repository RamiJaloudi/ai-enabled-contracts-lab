from __future__ import annotations

import re
from io import BytesIO

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from contractlab.legal_ai.contract_artifacts import (
    _add_page_number,
    _configure_styles,
    _paragraph_bottom_border,
    _scrub_metadata,
    _set_font,
)


def _safe_stem(title: str) -> str:
    value = re.sub(r"[^A-Za-z0-9._-]+", "-", title).strip("-._").lower()
    return value[:80] or "contract-chat-response"


def response_text_bytes(title: str, content: str, sources: list[str] | None = None) -> bytes:
    sections = [title.strip(), "", content.strip()]
    if sources:
        sections.extend(["", "Sources in this chat", *[f"- {source}" for source in sources]])
    sections.extend(["", "Drafting and educational assistance; not legal advice."])
    return "\n".join(sections).encode("utf-8")


def build_chat_response_docx(
    title: str,
    content: str,
    sources: list[str] | None = None,
) -> bytes:
    """Build a clean Word export using the standard business brief design preset."""
    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION.NEW_PAGE
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    _configure_styles(document)

    header = section.header.paragraphs[0]
    run = header.add_run("AI-ENABLED CONTRACTS LAB  |  CONTRACT CHAT")
    _set_font(run, size=8.5, color="666666", bold=True)
    _paragraph_bottom_border(header, color="D7DBE2", size=6)
    _add_page_number(section.footer.paragraphs[0])

    kicker = document.add_paragraph()
    kicker.paragraph_format.space_after = Pt(3)
    _set_font(kicker.add_run("CHAT RESPONSE"), size=9, color="2E74B5", bold=True)
    heading = document.add_paragraph()
    heading.paragraph_format.space_after = Pt(12)
    _set_font(heading.add_run(title.strip()), size=24, color="0B2545", bold=True)
    _paragraph_bottom_border(heading, color="2E74B5", size=12)

    for raw_line in content.splitlines():
        line = raw_line.strip()
        if not line:
            document.add_paragraph()
            continue
        markdown_heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if markdown_heading:
            level = min(len(markdown_heading.group(1)), 3)
            document.add_paragraph(markdown_heading.group(2), style=f"Heading {level}")
        elif re.match(r"^[-*]\s+", line):
            document.add_paragraph(re.sub(r"^[-*]\s+", "", line), style="List Bullet")
        elif re.match(r"^\d+[.)]\s+", line):
            document.add_paragraph(re.sub(r"^\d+[.)]\s+", "", line), style="List Number")
        else:
            document.add_paragraph(line)

    if sources:
        document.add_paragraph("Source documents", style="Heading 2")
        for source in sources:
            document.add_paragraph(source, style="List Bullet")

    notice = document.add_paragraph()
    notice.paragraph_format.space_before = Pt(12)
    notice.alignment = WD_ALIGN_PARAGRAPH.LEFT
    _paragraph_bottom_border(notice, color="D7DBE2", size=6)
    _set_font(
        notice.add_run("Drafting and educational assistance; not legal advice."),
        size=9,
        color="666666",
    )
    properties = document.core_properties
    properties.title = title.strip()
    properties.subject = "Contract chat response"
    properties.author = ""
    properties.last_modified_by = ""
    properties.keywords = ""
    buffer = BytesIO()
    document.save(buffer)
    return _scrub_metadata(buffer.getvalue())


def response_artifact_stem(title: str) -> str:
    return _safe_stem(title)
