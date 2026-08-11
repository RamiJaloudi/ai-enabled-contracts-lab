"""Build the AI-Enabled Contracts Lab Word guide from its Markdown source."""

from __future__ import annotations

import argparse
import json
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import (
    WD_ALIGN_PARAGRAPH,
    WD_BREAK,
    WD_LINE_SPACING,
    WD_TAB_ALIGNMENT,
    WD_TAB_LEADER,
)
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F4F6F7"
MID_GRAY = "667085"
WHITE = "FFFFFF"


@dataclass(slots=True)
class Block:
    kind: str
    value: object
    level: int = 0
    info: str = ""


def parse_markdown(path: Path) -> list[Block]:
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[Block] = []
    index = 0

    def starts_special(line: str, next_line: str = "") -> bool:
        stripped = line.strip()
        return (
            not stripped
            or stripped.startswith(("#", chr(96) * 3, ">", "- ", "* "))
            or stripped == "<!-- pagebreak -->"
            or re.match(r"^\d+\.\s+", stripped) is not None
            or (stripped.startswith("|") and next_line.strip().startswith("|"))
        )

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if not stripped:
            index += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", stripped)
        if heading:
            blocks.append(Block("heading", heading.group(2).strip(), len(heading.group(1))))
            index += 1
            continue

        if stripped == "<!-- pagebreak -->":
            blocks.append(Block("pagebreak", ""))
            index += 1
            continue

        if stripped.startswith(chr(96) * 3):
            language = stripped[3:].strip()
            index += 1
            code: list[str] = []
            while index < len(lines) and not lines[index].strip().startswith(chr(96) * 3):
                code.append(lines[index])
                index += 1
            index += 1
            blocks.append(Block("code", "\n".join(code), info=language))
            continue

        if stripped.startswith("|") and index + 1 < len(lines):
            separator = lines[index + 1].strip()
            cleaned = separator.replace("|", "").replace("-", "").replace(":", "").replace(" ", "")
            if separator.startswith("|") and not cleaned:
                rows: list[list[str]] = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    row = [cell.strip() for cell in lines[index].strip().strip("|").split("|")]
                    rows.append(row)
                    index += 1
                if len(rows) >= 2:
                    rows.pop(1)
                    blocks.append(Block("table", rows))
                continue

        if stripped.startswith(("- ", "* ")):
            items: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(("- ", "* ")):
                items.append(lines[index].strip()[2:].strip())
                index += 1
            blocks.append(Block("bullet", items))
            continue

        if re.match(r"^\d+\.\s+", stripped):
            items = []
            while index < len(lines):
                match = re.match(r"^\d+\.\s+(.+)$", lines[index].strip())
                if match is None:
                    break
                items.append(match.group(1).strip())
                index += 1
            blocks.append(Block("number", items))
            continue

        if stripped.startswith(">"):
            note_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                note_lines.append(lines[index].strip()[1:].strip())
                index += 1
            blocks.append(Block("note", " ".join(note_lines)))
            continue

        paragraph = [stripped]
        index += 1
        while index < len(lines):
            next_line = lines[index + 1] if index + 1 < len(lines) else ""
            if starts_special(lines[index], next_line):
                break
            paragraph.append(lines[index].strip())
            index += 1
        blocks.append(Block("paragraph", " ".join(paragraph)))

    return blocks


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(
    cell,
    top: int = 80,
    start: int = 100,
    bottom: int = 80,
    end: int = 100,
) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:tblHeader")) is None:
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        tr_pr.append(header)


def set_fixed_table_layout(table) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_internal_link(paragraph, text: str, anchor: str) -> None:
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_external_link(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    run_properties = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BLUE)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((color, underline))
    run.append(run_properties)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_PATTERN = re.compile(r"(\*\*.+?\*\*|\x60.+?\x60|\[[^\]]+\]\([^)]+\))")


def add_inline(paragraph, text: str) -> None:
    cursor = 0
    for match in INLINE_PATTERN.finditer(text):
        if match.start() > cursor:
            paragraph.add_run(text[cursor : match.start()])
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith(chr(96)):
            run = paragraph.add_run(token[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        else:
            link = re.match(r"\[([^\]]+)\]\(([^)]+)\)", token)
            if link:
                add_external_link(paragraph, link.group(1), link.group(2))
        cursor = match.end()
    if cursor < len(text):
        paragraph.add_run(text[cursor:])


def add_field(paragraph, instruction: str) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instr, separate, text, end))


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string("202124")
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.25

    heading_specs = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    styles["Heading 1"].paragraph_format.page_break_before = False

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)
        style.paragraph_format.left_indent = Inches(0.28)
        style.paragraph_format.first_line_indent = Inches(-0.18)
        style.paragraph_format.space_after = Pt(3)

    code_style = (
        styles["Code Block"]
        if "Code Block" in styles
        else styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    )
    code_style.font.name = "Consolas"
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = RGBColor.from_string("172B4D")
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.right_indent = Inches(0.18)
    code_style.paragraph_format.space_before = Pt(4)
    code_style.paragraph_format.space_after = Pt(6)
    code_style.paragraph_format.line_spacing = 1.0

    note_style = (
        styles["Guide Note"]
        if "Guide Note" in styles
        else styles.add_style("Guide Note", WD_STYLE_TYPE.PARAGRAPH)
    )
    note_style.font.name = "Calibri"
    note_style.font.size = Pt(10)
    note_style.font.italic = True
    note_style.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    note_style.paragraph_format.left_indent = Inches(0.22)
    note_style.paragraph_format.right_indent = Inches(0.12)
    note_style.paragraph_format.space_before = Pt(5)
    note_style.paragraph_format.space_after = Pt(8)

    toc1 = (
        styles["TOC Entry 1"]
        if "TOC Entry 1" in styles
        else styles.add_style("TOC Entry 1", WD_STYLE_TYPE.PARAGRAPH)
    )
    toc1.font.name = "Calibri"
    toc1.font.size = Pt(10)
    toc1.font.bold = True
    toc1.paragraph_format.space_after = Pt(3)

    toc2 = (
        styles["TOC Entry 2"]
        if "TOC Entry 2" in styles
        else styles.add_style("TOC Entry 2", WD_STYLE_TYPE.PARAGRAPH)
    )
    toc2.font.name = "Calibri"
    toc2.font.size = Pt(9)
    toc2.paragraph_format.left_indent = Inches(0.22)
    toc2.paragraph_format.space_after = Pt(1.5)


def configure_page(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("AI-ENABLED CONTRACTS LAB  |  USER & DEVELOPER GUIDE")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("AI-Enabled Contracts Lab 0.1.0   •   ")
    run.font.name = "Calibri"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)
    add_field(paragraph, "PAGE")


def add_cover(document: Document) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(70)

    accent = document.add_table(rows=1, cols=1)
    accent.alignment = WD_TABLE_ALIGNMENT.LEFT
    accent.autofit = False
    accent.columns[0].width = Inches(1.25)
    repeat_table_header(accent.rows[0])
    cell = accent.cell(0, 0)
    cell.width = Inches(1.25)
    set_cell_shading(cell, BLUE)
    cell.text = ""

    title = document.add_paragraph()
    title.paragraph_format.space_before = Pt(24)
    title.paragraph_format.space_after = Pt(10)
    run = title.add_run("AI-Enabled Contracts Lab")
    run.font.name = "Calibri"
    run.font.size = Pt(38)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("User & Developer Guide")
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string(BLUE)

    descriptor = document.add_paragraph()
    descriptor.paragraph_format.space_after = Pt(34)
    run = descriptor.add_run(
        "A private-first guide to the hybrid BERT and generative-LLM "
        "contract analysis, drafting, evaluation, and release workflow"
    )
    run.font.name = "Calibri"
    run.font.size = Pt(13)
    run.font.italic = True
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    audience = document.add_table(rows=2, cols=2)
    audience.alignment = WD_TABLE_ALIGNMENT.LEFT
    set_fixed_table_layout(audience)
    repeat_table_header(audience.rows[0])
    labels = (
        ("Audience", "Python learners, legal-AI researchers, and contributors"),
        ("Edition", "Version 0.1.0  •  August 2026"),
    )
    for row_index, (label, value) in enumerate(labels):
        row = audience.rows[row_index]
        row.cells[0].width = Inches(1.1)
        row.cells[1].width = Inches(5.2)
        for table_cell in row.cells:
            set_cell_margins(table_cell, 70, 90, 70, 90)
            set_cell_shading(table_cell, LIGHT_GRAY if row_index == 0 else WHITE)
        label_run = row.cells[0].paragraphs[0].add_run(label)
        label_run.bold = True
        label_run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
        value_run = row.cells[1].paragraphs[0].add_run(value)
        value_run.font.color.rgb = RGBColor.from_string("202124")

    disclaimer = document.add_paragraph()
    disclaimer.paragraph_format.space_before = Pt(55)
    disclaimer.paragraph_format.space_after = Pt(0)
    run = disclaimer.add_run("Educational software • Not legal advice • Human review required")
    run.font.name = "Calibri"
    run.font.size = Pt(10)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    document.add_page_break()


def heading_inventory(blocks: list[Block]) -> list[tuple[int, str, str]]:
    inventory: list[tuple[int, str, str]] = []
    counter = 1
    for block in blocks:
        if block.kind == "heading" and block.level in {2, 3}:
            inventory.append((block.level - 1, str(block.value), f"guide_sec_{counter:03d}"))
            counter += 1
    return inventory


def add_toc(
    document: Document,
    headings: list[tuple[int, str, str]],
    page_map: dict[str, int],
) -> None:
    title = document.add_paragraph()
    title.paragraph_format.space_after = Pt(5)
    run = title.add_run("Table of Contents")
    run.font.name = "Calibri"
    run.font.size = Pt(24)
    run.font.bold = True
    run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    intro = document.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    run = intro.add_run(
        "Click an entry in Word to navigate. Page numbering can vary with the "
        "installed fonts, Word version, printer settings, and accessibility options."
    )
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MID_GRAY)

    right_edge = Inches(6.45)
    for level, text, anchor in headings:
        style = "TOC Entry 1" if level == 1 else "TOC Entry 2"
        paragraph = document.add_paragraph(style=style)
        paragraph.paragraph_format.tab_stops.add_tab_stop(
            right_edge,
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )
        add_internal_link(paragraph, text, anchor)
        paragraph.add_run("\t")
        paragraph.add_run(str(page_map.get(text, "")))

    document.add_page_break()


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    set_fixed_table_layout(table)
    usable_width = 6.5
    if column_count == 2:
        widths = [2.0, 4.5]
    elif column_count == 3:
        widths = [1.55, 1.55, 3.4]
    else:
        widths = [usable_width / column_count] * column_count

    for row_index, values in enumerate(rows):
        row = table.rows[row_index]
        prevent_row_split(row)
        if row_index == 0:
            repeat_table_header(row)
        for column_index in range(column_count):
            cell = row.cells[column_index]
            cell.width = Inches(widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            fill = BLUE if row_index == 0 else (LIGHT_GRAY if row_index % 2 == 0 else WHITE)
            set_cell_shading(cell, fill)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            value = values[column_index] if column_index < len(values) else ""
            add_inline(paragraph, value)
            for run in paragraph.runs:
                run.font.name = "Calibri"
                run.font.size = Pt(8.5 if column_count >= 3 else 9)
                if row_index == 0:
                    run.font.bold = True
                    run.font.color.rgb = RGBColor.from_string(WHITE)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def shade_paragraph(paragraph, fill: str, border_color: str | None = None) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)
    if border_color:
        borders = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border_color)
        borders.append(left)
        properties.append(borders)


def create_numbering_instance(document: Document) -> int:
    numbering = document.part.numbering_part.element
    style_num_pr = document.styles["List Number"].element.pPr.numPr
    base_num_id = int(style_num_pr.numId.val)
    base_num = next(
        node
        for node in numbering.findall(qn("w:num"))
        if int(node.get(qn("w:numId"))) == base_num_id
    )
    abstract_num_id = base_num.find(qn("w:abstractNumId")).get(qn("w:val"))
    existing_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    new_num_id = max(existing_ids, default=0) + 1

    new_num = OxmlElement("w:num")
    new_num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_num_id)
    new_num.append(abstract)
    level_override = OxmlElement("w:lvlOverride")
    level_override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), "1")
    level_override.append(start_override)
    new_num.append(level_override)
    numbering.append(new_num)
    return new_num_id


def apply_numbering(paragraph, num_id: int) -> None:
    properties = paragraph._p.get_or_add_pPr()
    num_pr = properties.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        properties.append(num_pr)
    level = OxmlElement("w:ilvl")
    level.set(qn("w:val"), "0")
    number = OxmlElement("w:numId")
    number.set(qn("w:val"), str(num_id))
    num_pr.append(level)
    num_pr.append(number)


def add_body(
    document: Document,
    blocks: list[Block],
    headings: list[tuple[int, str, str]],
) -> None:
    bookmark_lookup = {(level, text): anchor for level, text, anchor in headings}
    bookmark_id = 100
    for block in blocks:
        if block.kind == "heading":
            if block.level == 1:
                continue
            word_level = min(3, block.level - 1)
            paragraph = document.add_paragraph(style=f"Heading {word_level}")
            add_inline(paragraph, str(block.value))
            anchor = bookmark_lookup.get((word_level, str(block.value)))
            if anchor:
                add_bookmark(paragraph, anchor, bookmark_id)
                bookmark_id += 1
            continue

        if block.kind == "paragraph":
            paragraph = document.add_paragraph()
            add_inline(paragraph, str(block.value))
            continue

        if block.kind == "pagebreak":
            document.add_page_break()
            continue

        if block.kind == "note":
            paragraph = document.add_paragraph(style="Guide Note")
            add_inline(paragraph, f"Important: {block.value}")
            shade_paragraph(paragraph, LIGHT_BLUE, BLUE)
            continue

        if block.kind == "bullet":
            for item in block.value:
                paragraph = document.add_paragraph(style="List Bullet")
                add_inline(paragraph, str(item))
            continue

        if block.kind == "number":
            num_id = create_numbering_instance(document)
            for item in block.value:
                paragraph = document.add_paragraph(style="List Number")
                apply_numbering(paragraph, num_id)
                add_inline(paragraph, str(item))
            continue

        if block.kind == "code":
            paragraph = document.add_paragraph(style="Code Block")
            shade_paragraph(paragraph, LIGHT_GRAY, "C9D3DD")
            lines = str(block.value).splitlines() or [""]
            for line_index, line in enumerate(lines):
                if line_index:
                    paragraph.add_run().add_break(WD_BREAK.LINE)
                run = paragraph.add_run(line if line else " ")
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
            continue

        if block.kind == "table":
            add_table(document, block.value)


def configure_document_settings(document: Document) -> None:
    settings = document.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "false")


def build(source: Path, output: Path, page_map_path: Path | None) -> None:
    blocks = parse_markdown(source)
    headings = heading_inventory(blocks)
    page_map: dict[str, int] = {}
    if page_map_path and page_map_path.exists():
        page_map = json.loads(page_map_path.read_text(encoding="utf-8"))

    document = Document()
    configure_styles(document)
    configure_page(document)
    configure_document_settings(document)
    add_cover(document)
    add_toc(document, headings, page_map)
    add_body(document, blocks, headings)

    properties = document.core_properties
    properties.title = "AI-Enabled Contracts Lab User and Developer Guide"
    properties.subject = (
        "Using, evaluating, and privately releasing the hybrid AI-Enabled Contracts Lab workbench"
    )
    properties.author = "Rami Jaloudi"
    properties.last_modified_by = "Rami Jaloudi"
    properties.keywords = "AI-Enabled Contracts Lab, BERT, generative LLM, contract chat, PyTorch, Streamlit, legal AI"
    properties.comments = "Generated from docs/AI_Enabled_Contracts_Lab_User_Guide.md"

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"[OK] wrote {output}")


def main() -> None:
    parser = argparse.ArgumentParser(description="Build the AI-Enabled Contracts Lab Word guide")
    parser.add_argument(
        "--source",
        type=Path,
        default=Path("docs/AI_Enabled_Contracts_Lab_User_Guide.md"),
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/AI_Enabled_Contracts_Lab_User_Guide.docx"),
    )
    parser.add_argument("--page-map", type=Path, default=None)
    args = parser.parse_args()
    build(args.source, args.output, args.page_map)


if __name__ == "__main__":
    main()
